#!/usr/bin/env python3
# Markdown → Word(.docx) 转换器（为本项目技术文档定制，纯 Python，不依赖 pandoc）
# 支持：# 标题 / 段落 / **加粗** / `行内代码` / 表格 / - 列表 / 1. 列表 / > 引用 / ``` 代码块 / --- 分割线
#
# 用法: python3 scripts/md2docx.py docs/TideWatch_技术文档.md docs/TideWatch_技术文档.docx

import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_runs(paragraph, text, bold_default=False):
    """解析 **加粗** 与 `行内代码`，写入 runs"""
    parts = re.split(r'(\*\*.+?\*\*|`[^`]+`)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            r.font.name = 'Menlo'
            r.font.size = Pt(9.5)
        else:
            r = paragraph.add_run(part)
        if bold_default:
            r.bold = True


def is_table_sep(line):
    return bool(re.match(r'^\s*\|[\s\-:|]+\|\s*$', line))


def split_row(line):
    cells = line.strip().strip('|').split('|')
    return [c.strip() for c in cells]


def convert(src_path, dst_path):
    with open(src_path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    i = 0
    while i < len(lines):
        line = lines[i]

        # 代码块
        if line.strip().startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过结尾 ```
            for cl in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(cl if cl else ' ')
                r.font.name = 'Menlo'
                r.font.size = Pt(9)
            continue

        # 表格
        if line.strip().startswith('|') and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            header = split_row(line)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            try:
                table.style = 'Table Grid'
            except KeyError:
                pass
            for j, h in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = ''
                add_runs(cell.paragraphs[0], h, bold_default=True)
            for row in rows:
                cells = table.add_row().cells
                for j, c in enumerate(row[:len(header)]):
                    cells[j].text = ''
                    add_runs(cells[j].paragraphs[0], c)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            level = len(m.group(1))
            doc.add_heading(re.sub(r'\*\*', '', m.group(2)), level=min(level, 4))
            i += 1
            continue

        # 分割线
        if re.match(r'^\s*---+\s*$', line):
            i += 1
            continue

        # 引用
        if line.strip().startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_runs(p, line.strip().lstrip('> ').strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # 无序列表
        if re.match(r'^\s*[-*]\s+', line):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^\s*[-*]\s+', '', line))
            i += 1
            continue

        # 有序列表
        if re.match(r'^\s*\d+\.\s+', line):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\s*\d+\.\s+', '', line))
            i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, line.strip())
        i += 1

    doc.save(dst_path)
    return dst_path


def main():
    if len(sys.argv) < 3:
        print('用法: python3 scripts/md2docx.py 输入.md 输出.docx')
        sys.exit(1)
    out = convert(sys.argv[1], sys.argv[2])
    print(f'已生成: {out}')


if __name__ == '__main__':
    main()
